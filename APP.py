import streamlit as st
import pandas as pd
from docx import Document

# ------------------------
# Funciones auxiliares
# ------------------------

def leer_excel(archivo_subido):
    try:
        df = pd.read_excel(archivo_subido)
        # Creamos un ID interno por fila para usar más adelante
        if "ID_INTERNO" not in df.columns:
            df.insert(0, "ID_INTERNO", range(1, len(df) + 1))
        return df, None
    except Exception as e:
        return None, f"Error al leer el Excel: {e}"

def leer_word_como_parrafos(archivo_subido):
    try:
        doc = Document(archivo_subido)
        # Filtramos párrafos vacíos para que la vista previa sea más limpia
        parrafos = [p.text for p in doc.paragraphs if p.text.strip() != ""]
        return parrafos, None
    except Exception as e:
        return None, f"Error al leer el Word: {e}"

# ------------------------
# Configuración básica de la app
# ------------------------

st.set_page_config(
    page_title="Generador de documentos judiciales",
    layout="wide"
)

st.title("📄 Generador de documentos judiciales – Fase 1")
st.caption("Paso 1: Cargar base en Excel y plantilla en Word con previsualización básica.")

# Estado de sesión
if "df_base" not in st.session_state:
    st.session_state.df_base = None

if "parrafos_plantilla" not in st.session_state:
    st.session_state.parrafos_plantilla = None

col1, col2 = st.columns(2)

# ------------------------
# Panel izquierdo: cargar Excel
# ------------------------
with col1:
    st.subheader("① Cargar base de datos (Excel)")

    archivo_excel = st.file_uploader(
        "Sube la base en Excel (.xlsx):",
        type=["xlsx", "xls"],
        key="uploader_excel"
    )

    if archivo_excel is not None:
        df, error = leer_excel(archivo_excel)
        if error:
            st.error(error)
        else:
            st.session_state.df_base = df
            st.success(f"Base cargada correctamente. Registros: {len(df)}")

            st.markdown("**Vista previa de las primeras filas:**")
            st.dataframe(df.head(10))

            st.markdown("**Columnas detectadas en la base:**")
            st.write(list(df.columns))

# ------------------------
# Panel derecho: cargar Word
# ------------------------
with col2:
    st.subheader("② Cargar plantilla (Word)")

    archivo_word = st.file_uploader(
        "Sube la plantilla en Word (.docx):",
        type=["docx"],
        key="uploader_word"
    )

    if archivo_word is not None:
        parrafos, error = leer_word_como_parrafos(archivo_word)
        if error:
            st.error(error)
        else:
            st.session_state.parrafos_plantilla = parrafos
            st.success("Plantilla Word cargada correctamente.")

            st.markdown("**Vista previa de los párrafos (solo texto):**")
            for i, p in enumerate(parrafos, start=1):
                st.markdown(f"**{i}.** {p}")

# ------------------------
# Estado general
# ------------------------
st.markdown("---")
st.subheader("Estado general")

if st.session_state.df_base is not None:
    st.success("✅ Base de datos cargada.")
else:
    st.warning("⚠️ Aún no has cargado la base de datos (Excel).")

if st.session_state.parrafos_plantilla is not None:
    st.success("✅ Plantilla Word cargada.")
else:
    st.warning("⚠️ Aún no has cargado la plantilla (Word).")

st.info("En el siguiente paso vamos a empezar el marcado de campos (JUZGADO, DEMANDANTE, etc.) sobre esta plantilla.")
# ------------------------
# PASO 3: Marcado de campos ({{...}}) + previsualización
# ------------------------
import re
import pandas as pd  # ya lo tienes arriba, pero no pasa nada si se repite

st.markdown("---")
st.header("③ Marcado de campos en la plantilla y previsualización")

# Verificamos que ambos estén cargados
if st.session_state.df_base is None or st.session_state.parrafos_plantilla is None:
    st.warning("Carga primero la base de datos (Excel) y la plantilla (Word).")
    st.stop()

df = st.session_state.df_base
parrafos = st.session_state.parrafos_plantilla

# Diccionario global: nombre_placeholder -> columna_excel
if "mapeo_placeholders" not in st.session_state:
    st.session_state.mapeo_placeholders = {}

st.write(
    "Detectamos variables dentro del texto con el formato {{NOMBRE}}. "
    "Aquí puedes vincular cada variable a una columna de la base."
)

# Para también poder avisar si hay placeholders sin mapear
placeholders_detectados_global = set()

# ---- Marcado de variables por párrafo ----
for idx, p in enumerate(parrafos):
    # Buscar placeholders tipo {{ NOMBRE }} dentro del párrafo
    placeholders = re.findall(r"{{\s*([^}]+?)\s*}}", p)

    if not placeholders:
        continue

    placeholders_unicos = sorted(set(placeholders))
    placeholders_detectados_global.update(placeholders_unicos)

    with st.expander(f"Párrafo {idx+1}"):
        st.markdown("### Contenido del párrafo:")
        st.write(p)

        st.markdown("### Variables detectadas en este párrafo:")

        for ph in placeholders_unicos:
            # Valor actual si ya habíamos mapeado esta variable antes
            valor_actual = st.session_state.mapeo_placeholders.get(ph, "(No vincular)")

            opciones = ["(No vincular)"] + list(df.columns)

            # Determinar índice por defecto del selectbox
            if valor_actual in df.columns:
                index_default = opciones.index(valor_actual)
            else:
                index_default = 0

            col_select = st.selectbox(
                f"Vincular la variable '{{{{{ph}}}}}' a una columna de la base:",
                options=opciones,
                index=index_default,
                key=f"ph_{idx}_{ph}"
            )

            # Actualizar mapeo global
            if col_select != "(No vincular)":
                st.session_state.mapeo_placeholders[ph] = col_select
            else:
                # Si el usuario elige "No vincular", la quitamos del diccionario (si existía)
                if ph in st.session_state.mapeo_placeholders:
                    del st.session_state.mapeo_placeholders[ph]

st.markdown("### 📝 Resumen de variables vinculadas")

if st.session_state.mapeo_placeholders:
    st.write(st.session_state.mapeo_placeholders)
else:
    st.info("Aún no has vinculado ninguna variable {{...}} a columnas de la base.")

# Aviso de variables detectadas pero no mapeadas
no_mapeadas = placeholders_detectados_global.difference(st.session_state.mapeo_placeholders.keys())
if no_mapeadas:
    st.warning(
        f"Estas variables fueron detectadas en la plantilla pero no están vinculadas a ninguna columna: "
        f"{', '.join(sorted(no_mapeadas))}"
    )

# ------------------------
# Previsualización con una fila de ejemplo
# ------------------------
st.markdown("---")
st.subheader("👁️ Previsualización del documento con una fila de la base")

if not st.session_state.mapeo_placeholders:
    st.info("Primero vincula al menos una variable {{...}} a alguna columna para poder previsualizar.")
else:
    # Selector de fila de ejemplo
    total_filas = len(df)
    fila_idx = st.number_input(
        "Selecciona el número de fila de la base para usar como ejemplo (1 = primera fila):",
        min_value=1,
        max_value=total_filas,
        value=1,
        step=1
    )

    fila = df.iloc[fila_idx - 1]

    st.caption(f"Mostrando previsualización usando la fila {fila_idx} de {total_filas}.")

    # Generar previsualización de cada párrafo
    st.markdown("### Resultado previsualizado:")

    for i, p in enumerate(parrafos, start=1):
        texto = p

        # Reemplazar cada placeholder mapeado por el valor correspondiente de la fila
        for ph, col in st.session_state.mapeo_placeholders.items():
            if col in df.columns:
                valor = fila[col]
                if pd.isna(valor):
                    valor = ""
                # Reemplazamos cualquier variante {{ NOMBRE }} / {{NOMBRE}} / {{   NOMBRE   }}
                patron = r"{{\s*" + re.escape(ph) + r"\s*}}"
                texto = re.sub(patron, str(valor), texto)

        st.markdown(f"**Párrafo {i}:**")
        st.write(texto)

